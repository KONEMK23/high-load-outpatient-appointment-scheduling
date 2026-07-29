from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from docx import Document

from build_bilingual_guides import (
    QA_ANSWER_TRANSLATIONS,
    QA_QUESTION_TRANSLATIONS,
    QA_SLIDE_MAP,
    SCRIPT_EXACT_TRANSLATIONS,
    SCRIPT_SUGGESTED_TRANSLATIONS,
    get_paths,
)


def check(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def paragraph_texts(doc: Document) -> list[str]:
    return [p.text for p in doc.paragraphs]


def assert_subsequence(source_items: list[str], output_items: list[str], label: str) -> None:
    source_pos = 0
    for item in output_items:
        if source_pos < len(source_items) and item == source_items[source_pos]:
            source_pos += 1
    check(
        source_pos == len(source_items),
        f"{label} source paragraph {source_pos + 1} was changed, removed, or reordered",
    )


def table_texts(doc: Document) -> list[list[list[str]]]:
    return [
        [[cell.text for cell in row.cells] for row in table.rows]
        for table in doc.tables
    ]


def verify_package(path: Path) -> None:
    check(path.exists() and path.stat().st_size > 0, f"Missing output: {path}")
    with ZipFile(path) as package:
        check(package.testzip() is None, f"Corrupt ZIP member in {path.name}")
        names = set(package.namelist())
        check("[Content_Types].xml" in names, f"Missing content types in {path.name}")
        check("word/document.xml" in names, f"Missing document.xml in {path.name}")
        check("word/styles.xml" in names, f"Missing styles.xml in {path.name}")


def verify_qa(source_path: Path, output_path: Path) -> tuple[int, int, int]:
    source = Document(source_path)
    output = Document(output_path)
    source_text = paragraph_texts(source)
    output_text = paragraph_texts(output)

    mappings = [p.text for p in output.paragraphs if p.style.name == "PPT Mapping"]
    question_translations = [
        p.text for p in output.paragraphs if p.style.name == "Chinese Question"
    ]
    translations = [
        p.text for p in output.paragraphs if p.style.name == "Chinese Translation"
    ]
    expected_mappings = [f"对应PPT：{QA_SLIDE_MAP[i]}" for i in range(1, 23)]
    expected_question_translations = [
        QA_QUESTION_TRANSLATIONS[i] for i in range(1, 23)
    ]
    expected_translations = [QA_ANSWER_TRANSLATIONS[i] for i in range(1, 23)]
    check(
        question_translations == expected_question_translations,
        "Q&A Chinese question translations are incomplete or out of order",
    )
    check(mappings == expected_mappings, "Q&A PowerPoint mappings are incomplete or out of order")
    check(
        translations == expected_translations,
        "Q&A Chinese translations are incomplete or out of order",
    )

    for i in range(1, 23):
        heading_index = next(
            j for j, text in enumerate(output_text) if text.startswith(f"Q{i}. ")
        )
        check(
            output_text[heading_index + 1] == expected_question_translations[i - 1],
            f"Q{i} Chinese question is not directly below its English heading",
        )
        check(
            output_text[heading_index + 2] == expected_mappings[i - 1],
            f"Q{i} mapping is not directly below its heading",
        )
        answer_index = next(
            j
            for j in range(heading_index + 1, len(output_text))
            if output_text[j].startswith("Answer in English:")
        )
        check(
            output_text[answer_index + 1] == expected_translations[i - 1],
            f"Q{i} translation is not directly below its English answer",
        )

    assert_subsequence(source_text, output_text, "Q&A")
    check(table_texts(source) == table_texts(output), "Q&A tables changed")
    return len(question_translations), len(mappings), len(translations)


def verify_script(source_path: Path, output_path: Path) -> tuple[int, int]:
    source = Document(source_path)
    output = Document(output_path)
    source_text = paragraph_texts(source)
    output_text = paragraph_texts(output)
    translations = [
        p.text for p in output.paragraphs if p.style.name == "Chinese Translation"
    ]
    expected_translation_set = set(SCRIPT_EXACT_TRANSLATIONS.values()) | set(
        SCRIPT_SUGGESTED_TRANSLATIONS.values()
    )

    check(
        set(translations) == expected_translation_set,
        "Speaker-script translations are incomplete or contain unexpected text",
    )
    check(
        len(translations) == len(expected_translation_set),
        "Speaker-script translations contain duplicates",
    )
    slides = [
        p.text
        for p in output.paragraphs
        if p.style.name == "Heading 1" and p.text.startswith("Slide ")
    ]
    check(len(slides) == 8, f"Expected 8 speaker slides, found {len(slides)}")

    for english, chinese in SCRIPT_EXACT_TRANSLATIONS.items():
        english_index = output_text.index(english)
        check(
            output_text[english_index + 1] == chinese,
            f"Translation is not directly below: {english[:50]}",
        )
    for slide_number, chinese in SCRIPT_SUGGESTED_TRANSLATIONS.items():
        slide_index = next(
            i
            for i, text in enumerate(output_text)
            if text.startswith(f"Slide {slide_number} - ")
        )
        suggested_index = output_text.index("Suggested wording", slide_index)
        english_index = suggested_index + 1
        check(
            output_text[english_index + 1] == chinese,
            f"Slide {slide_number} full script translation is misplaced",
        )

    assert_subsequence(source_text, output_text, "Speaker script")
    check(table_texts(source) == table_texts(output), "Speaker-script tables changed")
    return len(slides), len(translations)


def main() -> None:
    paths = get_paths()
    verify_package(paths["qa_output"])
    verify_package(paths["script_output"])
    qa_questions, qa_mappings, qa_translations = verify_qa(
        paths["qa_source"], paths["qa_output"]
    )
    speaker_slides, speaker_translations = verify_script(
        paths["script_source"], paths["script_output"]
    )

    print(f"QA question translations: {qa_questions}")
    print(f"QA mappings: {qa_mappings}")
    print(f"QA translations: {qa_translations}")
    print(f"Speaker slides: {speaker_slides}")
    print(f"Speaker translations: {speaker_translations}")
    print("English preservation: PASS")
    print("DOCX package integrity: PASS")


if __name__ == "__main__":
    main()
