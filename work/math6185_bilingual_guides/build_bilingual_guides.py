from __future__ import annotations

import os
import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph


QA_SLIDE_MAP = {
    1: "Slide 2",
    2: "Slides 5 and 13",
    3: "Slides 5 and 10",
    4: "Slide 10",
    5: "Slide 4",
    6: "Slides 4 and 9",
    7: "Slide 9",
    8: "Slides 6 and 11",
    9: "Slide 11",
    10: "Slides 6 and 11",
    11: "Slides 8 and 11",
    12: "Slides 7 and 12",
    13: "Slide 12",
    14: "Slides 7 and 12",
    15: "Slide 7",
    16: "Slide 12",
    17: "Slide 12",
    18: "Slides 2 and 14",
    19: "Slide 8",
    20: "Slide 8",
    21: "Slide 13",
    22: "Slide 14",
}


QA_QUESTION_TRANSLATIONS = {
    1: "问题中文：核心的运营问题是什么？",
    2: "问题中文：为什么最早可用时段策略是一个有用的基准？",
    3: "问题中文：你是如何验证Q2基准结果的？",
    4: "问题中文：为什么选择服务1,000名病人作为预热期？",
    5: "问题中文：为什么每次重复实验在到达人数达到50,000时停止？",
    6: "问题中文：为什么使用250次独立重复实验？",
    7: "问题中文：这里的95%置信区间是什么意思？",
    8: "问题中文：Q3策略是如何运行的？",
    9: "问题中文：为什么选择alpha2等于0.35？",
    10: "问题中文：Q3对常规病人公平吗？",
    11: "问题中文：既然常规病人的等待时间增加了，为什么仍然推荐Q3？",
    12: "问题中文：Q4的随叫随到常规病人策略是如何运行的？",
    13: "问题中文：为什么选择eta等于0.2、L等于0、H等于3？",
    14: "问题中文：为什么Q4对紧急等待时间的改善很小？",
    15: "问题中文：97.911%的利用率是否意味着Q4更优？",
    16: "问题中文：为什么不选择H等于2，而选择H等于3？",
    17: "问题中文：什么是事件等价筛选？",
    18: "问题中文：文献如何真正影响模型，而不只是用来装饰报告？",
    19: "问题中文：模型的主要局限是什么？",
    20: "问题中文：如果继续开展研究，下一步会做什么？",
    21: "问题中文：这项工作中你的原创贡献是什么？",
    22: "问题中文：你是如何使用AI辅助的？",
}


QA_ANSWER_TRANSLATIONS = {
    1: "中文对照：该诊所的名义负荷为0.98，因此几乎没有闲置容量。要缩短紧急病人的等待时间，就必须在不同病人类别之间重新分配延误，或者更有效地利用灵活重排预约。",
    2: "中文对照：它简单、对两类病人保持中立，而且容易解释。由于它为每位病人安排最早的空闲时段，因此可以展示在引入优先机制之前，共享系统本身的表现。",
    3: "中文对照：我从逻辑一致性、独立重复实验的精度以及预热期敏感性三个方面验证了Q2基准。在合并排队规则下，紧急和常规病人的平均等待时间几乎相同；相对置信区间半宽为4.10%；不同预热长度对应的置信区间也相互重叠。",
    4: "中文对照：这是一个得到敏感性分析支持的中间选择。预热500、1,000和2,000名已服务病人得到的95%置信区间高度重叠，因此研究结论对所选预热值并不敏感。",
    5: "中文对照：较长的单次运行能够在预热期之后提供足够多的观测值，从而形成稳定的单次重复估计，同时计算量仍然可控。最终精度是通过不同重复实验之间的置信区间来判断的，而不是仅凭单次运行长度判断。",
    6: "中文对照：独立重复实验可以形成均值的经验抽样分布，并据此计算常规的95%置信区间。250次重复实验使每个入选策略都达到了预先设定的相对半宽目标。",
    7: "中文对照：如果反复进行多组相互独立的仿真实验，并使用相同方法构造区间，那么这些区间大约有95%会包含真实的稳态均值。它描述的是给定模型下的蒙特卡洛随机误差，而不是模型假设本身是否正确的不确定性。",
    8: "中文对照：该策略观察紧急病人的积压量，并按照round(alpha1 NU + alpha2 NU squared)推迟选定的常规预约。当alpha1为0、alpha2为0.35时，随着紧急拥堵加剧，保护力度以非线性方式增加。",
    9: "中文对照：在测试过的非线性保护设置中，这是总体等待时间95%置信区间上限仍低于5天的最强设置。更高的保护强度会使常规病人等待和总体表现越过可行边界。",
    10: "中文对照：如果采用较宽泛的公平性定义，Q3并不公平。它虽然满足总体约束，却把常规病人的平均等待时间提高到了9.053天。因此，在实际应用前，我建议增加针对常规病人的服务水平约束。",
    11: "中文对照：因为题目设定的目标是在总体平均等待不超过5天的前提下，优先降低紧急病人的平均等待。按照这一目标，Q3带来的紧急等待改善最大，而且仍然可行。我的推荐带有一个条件：必须增加常规病人的保障措施。",
    12: "中文对照：该策略先保护近期容量，然后利用每日结束事件扫描未使用时段。它会释放符合条件的时段，并把等待时间最长且符合条件的随叫随到常规病人提前安排。",
    13: "中文对照：筛选结果表明，在98%的负荷下，H=4会保护过多容量。在提交的候选方案中，eta=0.2且H=3的设置既满足总体约束，又在AnyLogic直接验证的结果中取得了最好的紧急病人平均等待表现。",
    14: "中文对照：系统的名义负荷已经达到0.98，因此被保护的时段中很少有真正保持空闲的。重新预约能够维持较高利用率，但当需求几乎等于容量时，它无法为紧急病人创造出很多额外容量。",
    15: "中文对照：不能。高利用率说明被保护的容量没有被大量浪费，但研究目标是在满足总体约束的前提下降低紧急等待。虽然Q4在两类病人之间更加均衡，但Q3在紧急等待方面明显更好。",
    16: "中文对照：事件等价筛选中的H=2方案总体均值更低、紧急均值也相近，但它不能与H=3的250次AnyLogic直接重复实验结果进行严格的直接比较。因此，我保留了经过直接验证的提交方案，没有声称H=2与H=3之间已经完成正式优劣检验。",
    17: "中文对照：它是一种计算筛选实现，目的是在最终AnyLogic直接实验之外复现预约事件逻辑。它用于缩小候选参数范围，但最终入选策略所报告的结果来自AnyLogic的直接验证。",
    18: "中文对照：文献影响了实际测试的机制，包括作为基准的合并预约、为紧急病人保护容量、依据积压量动态调整优先级，以及对释放容量进行灵活重排。文献也促使我把等待时间的重新分配视为核心权衡。",
    19: "中文对照：模型假设需求是平稳泊松过程，服务时间确定、所有工作日相同，并且没有爽约、取消、偏好、病情恶化或加班。参数搜索范围也是有限的，因此不能证明所得方案是全局最优。",
    20: "中文对照：下一步我会增加常规病人的服务水平约束，在AnyLogic中直接验证最有潜力的Q4筛选候选方案，使用共同随机数改进相近方案之间的比较，并测试一个把Q3积压信号与Q4受控释放机制结合起来的混合策略。",
    21: "中文对照：我的贡献是在同一个仿真框架下统一构建并比较三个策略，利用置信区间上限明确检查可行性，并指出Q3的收益来自等待时间的重新分配，而不是没有代价的效率提升。",
    22: "中文对照：AI辅助用于起草、格式整理和一致性检查。建模假设和策略由我选择，仿真实验由我运行和检查，数值输出及权衡也由我核对和评估；我对最终提交内容和演示承担责任。",
}


SCRIPT_EXACT_TRANSLATIONS = {
    "Structure: Slides 1-8 are the spoken core; slides 9-14 are backup for questions.":
        "中文对照：结构安排：Slides 1–8是正式讲解内容；Slides 9–14用于回答提问时提供补充证据。",
    "Timing: Aim for roughly 60-75 seconds on slides 1-4 and 45-60 seconds on slides 5-8.":
        "中文对照：时间安排：Slides 1–4每页大约讲60–75秒，Slides 5–8每页大约讲45–60秒。",
    "Short version: If time is limited, omit the grey optional lines rather than rushing the final recommendation.":
        "中文对照：精简版本：如果时间有限，应省略灰色的可选内容，不要因为赶时间而仓促讲完最终建议。",
    "Demonstration: Do not run AnyLogic live. Use the model-structure backup slide if implementation details are requested.":
        "中文对照：演示安排：不要现场运行AnyLogic。如果老师询问实现细节，使用模型结构备用页回答。",
    "Questions: Answer directly, give one result or mechanism, then state the limitation if it matters.":
        "中文对照：回答问题：先直接给出结论，再提供一个结果数字或机制；如果相关，再说明局限性。",
    "Purpose: Frame the decision and state the result immediately.":
        "中文对照：本页目的：明确决策问题，并立即给出主要结果。",
    "Purpose: Explain why the scheduling problem is structurally difficult and connect it to the literature.":
        "中文对照：本页目的：解释该预约问题为何在结构上很困难，并把问题与相关文献联系起来。",
    "Purpose: Show the answer-first comparison across all policies.":
        "中文对照：本页目的：先给结论，再比较所有策略的表现。",
    "Purpose: Explain the simulation structure, warm-up, stopping rule and precision.":
        "中文对照：本页目的：解释仿真结构、预热期、停止规则和统计精度。",
    "Purpose: Validate first-available scheduling and establish the baseline.":
        "中文对照：本页目的：验证最早可用时段策略，并建立比较基准。",
    "Purpose: Explain the workload-responsive rule, result and fairness cost.":
        "中文对照：本页目的：解释基于工作负荷的规则、所得结果以及公平性代价。",
    "Purpose: Explain the on-call routine mechanism and why its improvement is modest.":
        "中文对照：本页目的：解释随叫随到常规病人的机制，以及其改善幅度为何有限。",
    "Purpose: Resolve the decision and state limitations and next actions.":
        "中文对照：本页目的：给出最终选择，并说明局限性和下一步工作。",
    "[Optional if time is short] Omit the list of all three policy names and move directly from the decision question to the main result.":
        "中文对照：[时间不足时可省略] 不必逐一说出三个策略名称，可从决策问题直接转到主要结果。",
    "[Optional if time is short] Omit the detailed literature mechanism examples and retain only the central trade-off.":
        "中文对照：[时间不足时可省略] 省略文献中各机制的详细例子，只保留核心权衡。",
    "[Optional if time is short] Omit the Q4 routine and overall figures; retain the urgent comparison and the Q3 fairness warning.":
        "中文对照：[时间不足时可省略] 省略Q4的常规和总体数值，只保留紧急等待比较以及对Q3公平性的提醒。",
    "[Optional if time is short] Omit the final sentence listing all three relative half-widths.":
        "中文对照：[时间不足时可省略] 省略最后一句中三个相对置信区间半宽的具体数值。",
    "[Optional if time is short] Omit the three warm-up point estimates and state only that their confidence intervals overlap.":
        "中文对照：[时间不足时可省略] 省略三个预热期的点估计，只说明它们的置信区间相互重叠。",
    "[Optional if time is short] Omit the verbal expansion of the formula after showing it.":
        "中文对照：[时间不足时可省略] 展示公式后，不必再逐项口头展开公式。",
    "[Optional if time is short] Omit the conservative upper-confidence-limit value if time is short.":
        "中文对照：[时间不足时可省略] 可以不说保守检验所用的置信区间上限数值。",
    "[Optional if time is short] Omit the detailed end-of-day selection rule and retain only hold, release and rebook.":
        "中文对照：[时间不足时可省略] 省略每日结束时的详细选择规则，只保留“保护、释放、重新预约”三个步骤。",
    "[Optional if time is short] Omit the list of alternative routine guardrail formulations.":
        "中文对照：[时间不足时可省略] 省略常规病人保障措施的不同具体形式。",
    "[Optional if time is short] Omit the hybrid-policy next step.":
        "中文对照：[时间不足时可省略] 省略下一步测试混合策略的内容。",
}


SCRIPT_SUGGESTED_TRANSLATIONS = {
    1: "中文对照：大家上午好。本项目研究的是，在不让总体平均等待时间超过5天的情况下，一个高利用率门诊如何缩短紧急病人的等待时间。我在相同需求和容量条件下建立了三个策略模型：作为基准的最早可用时段策略、基于工作负荷的保护规则，以及随叫随到常规病人的重新预约策略。主要结果很明确：基于工作负荷的Q3策略带来了最大的紧急等待改善，将其从最早可用时段策略下的4.148天下降到2.123天，降幅为48.8%。Q3的总体均值为4.703天，因此其95%置信区间上限仍低于5天。不过，常规病人的等待时间上升到了9.053天。因此，针对题目给定的目标，我推荐Q3，但前提是增加针对常规病人的服务水平保障。",
    2: "中文对照：这个问题的困难并不是缺少一种预约逻辑，而是系统缺少闲置容量。紧急病人平均每天到达3.69人，常规病人平均每天到达2.19人，而每天只有6个预约时段，因此名义负荷为0.98。在这种情况下，如果不改变由谁等待或何时释放容量，任何策略都无法真正为紧急病人创造更多早期预约机会。预约调度文献用不同方式说明了同一点：容量预留、动态优先级和灵活重新预约都能保护紧急需求，但每一种保护机制都会带来机会成本。因此，我把问题视为一种权衡，而不是寻找一个可以同时改善所有病人类别的策略。评价目标是在总体平均等待时间不超过5天的前提下，最小化紧急病人的平均等待时间。",
    3: "中文对照：这张比较图是本研究的核心。在最早可用时段策略下，紧急、常规和总体平均等待时间都约为4.148天，因为该策略不区分病人类别。Q3显著改变了等待时间的分配：紧急病人的平均等待降至2.123天，而常规病人的等待升至9.053天，总体均值为4.703天。Q4更加均衡：紧急病人为4.009天，常规病人为4.574天，总体为4.219天。三个入选策略都满足总体不超过5天的约束。我推荐Q3，是因为题目目标优先考虑紧急病人的平均等待，而Q3将该指标改善了大约2天。这张图也清楚显示了服务风险：总体可行并不意味着常规病人也获得了良好服务。",
    4: "中文对照：三个模型采用相同的到达过程和容量假设，因此结果差异可以归因于策略逻辑。紧急和常规病人按照相互独立的泊松过程到达。每个策略在每天6个时段的预约日历中分配或调整预约，等待时间从病人到达时刻计算到预约时刻。我把每次重复实验中最初服务的1,000名病人从统计中排除作为预热期，之后继续运行，直到到达人数达到50,000。实验包含250次独立重复，并报告95%置信区间。精度目标是总体均值的相对置信区间半宽不超过5%。入选结果都满足这一目标：Q2为4.10%，Q3为4.04%，Q4为3.49%。这种设计同时处理了初始瞬态偏差和蒙特卡洛随机误差。",
    5: "中文对照：Q2采用最早可用时段策略。它的重要性在于展示了所有病人在没有类别保护的情况下竞争最早空闲时段时的系统表现。总体平均等待为4.148天，95%置信区间为3.978至4.318天。紧急和常规病人的均值几乎相同，这正符合合并使用最早可用规则时的预期。我还检查了预热期选择。预热500、1,000和2,000名已服务病人时，总体均值分别为4.133、4.148和4.172天，而且置信区间高度重叠。这不能证明完全不存在瞬态影响，但说明在这些合理的预热选择下，报告结果并不敏感。因此，我把Q2作为一个稳定且容易解释的基准。",
    6: "中文对照：Q3通过基于工作负荷的推迟规则保护紧急需求。当紧急积压量NU增加时，选定的常规预约会按照alpha1乘以NU再加上alpha2乘以NU平方、然后取整的数值向后移动。入选参数为alpha1等于0、alpha2等于0.35，因此保护力度会随着紧急拥堵以非线性方式增长。结果显示紧急病人的等待大幅改善：平均等待降至2.123天，比FAS基准低48.8%。总体均值为4.703天，其95%置信区间上限为4.893，因此即使采用保守的上限检验也仍然可行。代价同样清楚：常规病人的平均等待上升至9.053天。这不是隐藏的副作用，而是稀缺容量被重新分配的具体机制。",
    7: "中文对照：Q4利用随叫随到的常规病人来回收被释放的容量。当eta等于0.2、L等于0、H等于3时，该策略先保护近期时段，然后在每天结束时扫描未使用的时段。如果某个合适的时段被释放，就把等待时间最长且符合条件的随叫随到常规病人提前安排。这样可以保持容量利用：预约时段利用率为97.911%，被重新预约的病人平均节省0.515天。不过，紧急病人的平均等待仍为4.009天，只比FAS低0.139天，即3.3%。常规和总体均值分别为4.574天和4.219天。这说明灵活重新预约能够有效利用被释放的容量，但在名义负荷为0.98时，真正可以释放的闲置容量很少。Q4可行且较为均衡，但在题目设定的紧急等待目标上不如Q3。",
    8: "中文对照：我推荐Q3，因为它最符合题目给定的目标：在总体平均等待不超过5天的条件下，最小化紧急病人的平均等待。Q3的紧急均值为2.123天，总体95%置信区间上限为4.893天。该决策的统计精度也足够，相对半宽为4.04%。但我不会建议原样直接实施这一规则。常规病人的平均等待达到9.053天，说明总体约束过于宽松，因此实际运营要求还应增加常规病人的平均等待上限、百分位服务目标或明确的类别权重。模型还假设需求平稳、服务时间确定，并且没有爽约、取消、病情恶化或加班。下一项有价值的实验，是把Q3的工作负荷信号与Q4式的受控释放机制结合起来形成混合策略。",
}


def get_paths() -> dict[str, Path]:
    project_root = Path(
        os.environ.get(
            "MATH6185_PROJECT_ROOT",
            Path(__file__).resolve().parents[2],
        )
    ).resolve()
    if not (project_root / "outputs").exists():
        project_root = project_root.parents[1]
    output_dir = project_root / "outputs" / "zw1f25_MATH6185_presentation"
    return {
        "project_root": project_root,
        "output_dir": output_dir,
        "qa_source": output_dir / "zw1f25_MATH6185_QA_Guide.docx",
        "script_source": output_dir / "zw1f25_MATH6185_Speaker_Script.docx",
        "qa_output": output_dir / "zw1f25_MATH6185_QA_Guide_Bilingual.docx",
        "script_output": output_dir / "zw1f25_MATH6185_Speaker_Script_Bilingual.docx",
    }


def ensure_styles(doc: Document) -> None:
    if "Chinese Question" not in doc.styles:
        style = doc.styles.add_style("Chinese Question", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(10)
        style.font.bold = True
        style.font.color.rgb = RGBColor(55, 65, 81)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.keep_with_next = True
        style.element.get_or_add_rPr().get_or_add_rFonts().set(
            qn("w:eastAsia"), "Microsoft YaHei"
        )

    if "PPT Mapping" not in doc.styles:
        style = doc.styles.add_style("PPT Mapping", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(9)
        style.font.bold = True
        style.font.color.rgb = RGBColor(31, 111, 139)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True
        style.element.get_or_add_rPr().get_or_add_rFonts().set(
            qn("w:eastAsia"), "Microsoft YaHei"
        )

    if "Chinese Translation" not in doc.styles:
        style = doc.styles.add_style("Chinese Translation", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(9)
        style.font.color.rgb = RGBColor(55, 65, 81)
        style.paragraph_format.space_before = Pt(1)
        style.paragraph_format.space_after = Pt(5)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(
            qn("w:eastAsia"), "Microsoft YaHei"
        )


def insert_after(paragraph: Paragraph, text: str, style: str) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    inserted = Paragraph(new_p, paragraph._parent)
    inserted.style = style
    run = inserted.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(
        qn("w:eastAsia"), "Microsoft YaHei"
    )
    return inserted


def build_qa_guide(source: Path, target: Path) -> None:
    doc = Document(source)
    ensure_styles(doc)
    questions_seen: set[int] = set()
    answers_seen: set[int] = set()
    current_question: int | None = None

    for paragraph in list(doc.paragraphs):
        heading_match = re.match(r"Q(\d+)\.\s", paragraph.text)
        if paragraph.style.name == "Heading 1" and heading_match:
            current_question = int(heading_match.group(1))
            questions_seen.add(current_question)
            chinese_question = insert_after(
                paragraph, QA_QUESTION_TRANSLATIONS[current_question], "Chinese Question"
            )
            insert_after(
                chinese_question,
                f"对应PPT：{QA_SLIDE_MAP[current_question]}",
                "PPT Mapping",
            )
            continue

        if paragraph.text.startswith("Answer in English:"):
            if current_question is None:
                raise ValueError("English answer found before a numbered question")
            insert_after(
                paragraph,
                QA_ANSWER_TRANSLATIONS[current_question],
                "Chinese Translation",
            )
            answers_seen.add(current_question)

    expected = set(range(1, 23))
    if questions_seen != expected or answers_seen != expected:
        raise ValueError(
            f"Q&A anchors incomplete: headings={sorted(questions_seen)}, "
            f"answers={sorted(answers_seen)}"
        )

    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)


def build_speaker_script(source: Path, target: Path) -> None:
    doc = Document(source)
    ensure_styles(doc)
    current_slide: int | None = None
    expect_suggested_wording = False
    suggested_seen: set[int] = set()

    for paragraph in list(doc.paragraphs):
        slide_match = re.match(r"Slide\s+(\d+)\s+-\s", paragraph.text)
        if paragraph.style.name == "Heading 1" and slide_match:
            current_slide = int(slide_match.group(1))
            expect_suggested_wording = False
            continue

        if paragraph.text == "Suggested wording":
            expect_suggested_wording = True
            continue

        if expect_suggested_wording and paragraph.text.strip():
            if current_slide not in SCRIPT_SUGGESTED_TRANSLATIONS:
                raise ValueError(f"Missing suggested-wording translation for Slide {current_slide}")
            insert_after(
                paragraph,
                SCRIPT_SUGGESTED_TRANSLATIONS[current_slide],
                "Chinese Translation",
            )
            suggested_seen.add(current_slide)
            expect_suggested_wording = False

        if paragraph.text in SCRIPT_EXACT_TRANSLATIONS:
            insert_after(
                paragraph,
                SCRIPT_EXACT_TRANSLATIONS[paragraph.text],
                "Chinese Translation",
            )

    if suggested_seen != set(range(1, 9)):
        raise ValueError(f"Suggested wording incomplete: {sorted(suggested_seen)}")

    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)


def main() -> None:
    paths = get_paths()
    build_qa_guide(paths["qa_source"], paths["qa_output"])
    build_speaker_script(paths["script_source"], paths["script_output"])
    print(paths["qa_output"])
    print(paths["script_output"])


if __name__ == "__main__":
    main()
