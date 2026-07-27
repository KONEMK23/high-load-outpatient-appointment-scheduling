from pathlib import Path
import os
import sys

from docx import Document


HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from build_bilingual_guides import (  # noqa: E402
    QA_ANSWER_TRANSLATIONS,
    QA_SLIDE_MAP,
    SCRIPT_EXACT_TRANSLATIONS,
    get_paths,
)


def check(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def main() -> None:
    paths = get_paths()

    check(set(QA_SLIDE_MAP) == set(range(1, 23)), "Slide map must cover Q1-Q22")
    check(
        set(QA_ANSWER_TRANSLATIONS) == set(range(1, 23)),
        "Answer translations must cover Q1-Q22",
    )
    check(
        all(text.strip() and text.startswith("中文对照：") for text in QA_ANSWER_TRANSLATIONS.values()),
        "Every Q&A translation must be non-empty and use the Chinese label",
    )

    qa = Document(paths["qa_source"])
    script = Document(paths["script_source"])
    qa_headings = [
        p.text
        for p in qa.paragraphs
        if p.style.name == "Heading 1" and p.text.startswith("Q") and ". " in p.text
    ]
    qa_answers = [p.text for p in qa.paragraphs if p.text.startswith("Answer in English:")]
    script_slides = [
        p.text
        for p in script.paragraphs
        if p.style.name == "Heading 1" and p.text.startswith("Slide ") and " - " in p.text
    ]
    script_wording = [p.text for p in script.paragraphs if p.text == "Suggested wording"]

    check(len(qa_headings) == 22, f"Expected 22 Q&A headings, found {len(qa_headings)}")
    check(len(qa_answers) == 22, f"Expected 22 English answers, found {len(qa_answers)}")
    check(len(script_slides) == 8, f"Expected 8 slide sections, found {len(script_slides)}")
    check(len(script_wording) == 8, f"Expected 8 wording headings, found {len(script_wording)}")

    missing_script = [
        p.text
        for p in script.paragraphs
        if (
            p.text.startswith(("Structure:", "Timing:", "Short version:", "Demonstration:", "Questions:", "Purpose:"))
            or p.text.startswith("[Optional if time is short]")
        )
        and p.text not in SCRIPT_EXACT_TRANSLATIONS
    ]
    check(not missing_script, f"Missing exact speaker translations: {missing_script}")

    print("Coverage tests: PASS")
    print(f"Q&A questions: {len(qa_headings)}")
    print(f"Speaker slides: {len(script_slides)}")


if __name__ == "__main__":
    main()
