from __future__ import annotations

import hashlib
import os
import posixpath
from collections import Counter
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree
from PIL import Image, ImageChops
from pypdf import PdfReader


NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
}

SOURCE_HASHES = {
    "zw1f25_MATH6185_Presentation.pptx": (
        "3BC80AB3644ABF37C1F21E930BF7B07D0E78DA0A8F3F54CE2A54C75D2EED9E19"
    ),
    "zw1f25_MATH6185_Presentation.pdf": (
        "C76D870BE27359AFFC5976BDB757A5B4C214D56DED4CA6557074E8955D34EAAA"
    ),
    "zw1f25_MATH6185_QA_Guide_Bilingual.docx": (
        "F9DDBCDB1659499ECC503D3C8F01ECA054B9EDA95C49BDDE283BDA7EA0A890B4"
    ),
}


def project_root() -> Path:
    configured = os.environ.get("MATH6185_PROJECT_ROOT")
    if configured:
        return Path(configured).resolve()
    candidate = Path(__file__).resolve().parents[2]
    if (candidate / "outputs").exists():
        return candidate
    return candidate.parents[1]


def output_dir() -> Path:
    return project_root() / "outputs" / "zw1f25_MATH6185_presentation"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def xml(zf: ZipFile, part: str):
    return etree.fromstring(zf.read(part))


def rels_part(part: str) -> str:
    folder, name = posixpath.split(part)
    return posixpath.join(folder, "_rels", f"{name}.rels")


def resolve_target(source_part: str, target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(source_part), target))


def relationships(zf: ZipFile, part: str) -> dict[str, dict[str, str]]:
    rel_part = rels_part(part)
    if rel_part not in zf.namelist():
        return {}
    result = {}
    root = xml(zf, rel_part)
    for rel in root.xpath("//pr:Relationship", namespaces=NS):
        result[rel.get("Id")] = {
            "type": rel.get("Type", "").rsplit("/", 1)[-1],
            "target": resolve_target(part, rel.get("Target", "")),
            "mode": rel.get("TargetMode", "Internal"),
        }
    return result


def ordered_slide_parts(zf: ZipFile) -> list[str]:
    presentation_part = "ppt/presentation.xml"
    root = xml(zf, presentation_part)
    rels = relationships(zf, presentation_part)
    parts = []
    for slide_id in root.xpath("//p:sldIdLst/p:sldId", namespaces=NS):
        rid = slide_id.get(f"{{{NS['r']}}}id")
        parts.append(rels[rid]["target"])
    return parts


def text_of(zf: ZipFile, part: str) -> list[str]:
    root = xml(zf, part)
    return root.xpath("//a:t/text()", namespaces=NS)


def note_part(zf: ZipFile, slide_part: str) -> str:
    for rel in relationships(zf, slide_part).values():
        if rel["type"] == "notesSlide":
            return rel["target"]
    raise AssertionError(f"Missing notes relationship for {slide_part}")


def normalized_relationship_inventory(zf: ZipFile, slide_part: str) -> Counter:
    inventory = Counter()
    for rel in relationships(zf, slide_part).values():
        if rel["mode"] != "Internal":
            inventory[(rel["type"], "external")] += 1
            continue
        suffix = Path(rel["target"]).suffix.lower()
        inventory[(rel["type"], suffix)] += 1
    return inventory


def slide_size(zf: ZipFile) -> tuple[str, str]:
    root = xml(zf, "ppt/presentation.xml")
    size = root.find("p:sldSz", namespaces=NS)
    assert size is not None
    return size.get("cx"), size.get("cy")


def package_counts(zf: ZipFile) -> Counter:
    counts = Counter()
    for name in zf.namelist():
        if name.startswith("ppt/slideMasters/") and name.endswith(".xml"):
            counts["masters"] += 1
        elif name.startswith("ppt/slideLayouts/") and name.endswith(".xml"):
            counts["layouts"] += 1
        elif name.startswith("ppt/theme/") and name.endswith(".xml"):
            counts["themes"] += 1
    return counts


def validate_pptx(source: Path, target: Path) -> None:
    with ZipFile(source) as source_zip, ZipFile(target) as target_zip:
        assert source_zip.testzip() is None
        assert target_zip.testzip() is None
        source_slides = ordered_slide_parts(source_zip)
        target_slides = ordered_slide_parts(target_zip)
        assert len(source_slides) == 14
        assert len(target_slides) == 8
        assert slide_size(source_zip) == slide_size(target_zip)
        assert package_counts(source_zip) == package_counts(target_zip)

        for index, (source_slide, target_slide) in enumerate(
            zip(source_slides[:8], target_slides), start=1
        ):
            assert text_of(source_zip, source_slide) == text_of(
                target_zip, target_slide
            ), f"Slide {index} text changed"
            assert normalized_relationship_inventory(
                source_zip, source_slide
            ) == normalized_relationship_inventory(
                target_zip, target_slide
            ), f"Slide {index} relationship inventory changed"
            source_note = note_part(source_zip, source_slide)
            target_note = note_part(target_zip, target_slide)
            assert text_of(source_zip, source_note) == text_of(
                target_zip, target_note
            ), f"Slide {index} speaker notes changed"

    print("PPTX slide count and source preservation: PASS")


def validate_pdf(source: Path, target: Path) -> None:
    source_pdf = PdfReader(str(source))
    target_pdf = PdfReader(str(target))
    assert len(source_pdf.pages) == 14
    assert len(target_pdf.pages) == 8
    for index, (source_page, target_page) in enumerate(
        zip(source_pdf.pages[:8], target_pdf.pages), start=1
    ):
        assert source_page.mediabox == target_page.mediabox, index
        assert source_page.cropbox == target_page.cropbox, index
    print("PDF page count and geometry: PASS")


def validate_docx(target: Path) -> None:
    document = Document(target)
    question_headings = [
        p.text
        for p in document.paragraphs
        if p.style.name == "Heading 2" and p.text.startswith("Q")
    ]
    slide_headings = [
        p.text
        for p in document.paragraphs
        if p.style.name == "Heading 1" and p.text.startswith("Slide ")
    ]
    assert len(question_headings) == 18
    assert len(slide_headings) == 8
    assert len(document.tables) == 1
    assert [row.cells[0].text for row in document.tables[0].rows[1:]] == [
        f"Slide {index}" for index in range(1, 9)
    ]

    all_text = "\n".join(p.text for p in document.paragraphs)
    all_text += "\n" + "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    for token in (
        "0.98",
        "250",
        "50,000",
        "1,000",
        "4.148",
        "2.123",
        "9.053",
        "4.703",
        "4.009",
        "4.574",
        "4.219",
        "97.911%",
        "0.515",
        "4.893",
    ):
        assert token in all_text, f"Missing DOCX numerical anchor: {token}"

    for phrase in ("event-equivalent", "H = 2", "AI assistance"):
        assert phrase.lower() not in all_text.lower(), f"Backup-only topic: {phrase}"
    print("DOCX structure and numerical anchors: PASS")


def compare_render_directories(source_dir: Path, target_dir: Path, label: str) -> None:
    source_images = sorted(source_dir.glob("*.png"))
    target_images = sorted(target_dir.glob("*.png"))
    assert len(source_images) >= 8 and len(target_images) == 8, (
        label,
        len(source_images),
        len(target_images),
    )
    for index, (source_path, target_path) in enumerate(
        zip(source_images[:8], target_images), start=1
    ):
        source_image = Image.open(source_path).convert("RGBA")
        target_image = Image.open(target_path).convert("RGBA")
        assert source_image.size == target_image.size, (label, index)
        assert ImageChops.difference(source_image, target_image).getbbox() is None, (
            label,
            index,
        )
    print(f"{label} render identity: PASS")


def validate_source_hashes(out: Path) -> None:
    for name, expected in SOURCE_HASHES.items():
        actual = sha256(out / name)
        assert actual == expected, f"Source changed: {name} {actual}"
    print("Original source hashes unchanged: PASS")


def main() -> None:
    out = output_dir()
    source_pptx = out / "zw1f25_MATH6185_Presentation.pptx"
    target_pptx = out / "zw1f25_MATH6185_Presentation_8_Slides.pptx"
    source_pdf = out / "zw1f25_MATH6185_Presentation.pdf"
    target_pdf = out / "zw1f25_MATH6185_Presentation_8_Slides.pdf"
    target_docx = out / "zw1f25_MATH6185_QA_Guide_PPT_Focused_Bilingual.docx"

    for path in (target_pptx, target_pdf, target_docx):
        assert path.exists() and path.stat().st_size > 20_000, path

    validate_pptx(source_pptx, target_pptx)
    validate_pdf(source_pdf, target_pdf)
    validate_docx(target_docx)

    tmp = Path(__file__).resolve().parent / "tmp"
    compare_render_directories(
        tmp / "template-inspect" / "source-slides",
        tmp / "final-artifact-render",
        "PPTX",
    )
    compare_render_directories(
        tmp / "pdf-source-render",
        tmp / "pdf-final-render",
        "PDF",
    )
    validate_source_hashes(out)
    print("FINAL EIGHT-SLIDE PACKAGE VALIDATION PASS")


if __name__ == "__main__":
    main()
