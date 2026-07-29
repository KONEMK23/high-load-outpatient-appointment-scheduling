from __future__ import annotations

import os
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


QUESTIONS = [
    {
        "slide": 1,
        "question_en": "What is the central decision problem?",
        "question_zh": "核心决策问题是什么？",
        "answer_en": (
            "The clinic is almost fully loaded, so the decision is how to reduce urgent "
            "patient waiting without letting the overall mean waiting time exceed five "
            "days. The three models test different ways of allocating the same scarce "
            "appointment capacity: neutral first-available scheduling, workload-responsive "
            "protection, and controlled release with rebooking."
        ),
        "answer_zh": (
            "该门诊几乎满负荷运行，因此核心决策是在不让总体平均等待时间超过5天的前提下，"
            "如何缩短紧急患者的等待时间。三个模型用同一份稀缺预约容量测试了三种分配方式："
            "中立的最早可用时段、随紧急积压变化的容量保护，以及受控释放并重新预约。"
        ),
        "keywords": "decision problem, urgent waiting, five-day constraint",
        "cue": "Mechanism cue: reallocate scarce capacity; do not create new capacity.",
    },
    {
        "slide": 1,
        "question_en": "What is the main result, and why is Q3 recommended?",
        "question_zh": "主要结果是什么，为什么推荐Q3？",
        "answer_en": (
            "Q3 gives the largest urgent-wait improvement: the urgent mean falls from "
            "4.148 days under Q2 to 2.123 days, a reduction of about 48.8%. Its overall "
            "mean is 4.703 days and the upper end of its 95% confidence interval is 4.893, "
            "so it remains within the five-day constraint. The recommendation is conditional "
            "because routine waiting rises to 9.053 days."
        ),
        "answer_zh": (
            "Q3带来的紧急等待改善最大：紧急患者平均等待从Q2的4.148天下降到2.123天，"
            "降幅约48.8%。Q3的总体均值为4.703天，95%置信区间上限为4.893天，仍满足"
            "5天约束。但这一推荐是有条件的，因为常规患者平均等待上升到了9.053天。"
        ),
        "keywords": "main result, recommend Q3, 48.8%, constraint",
        "cue": "Number cue: 2.123 urgent; 4.893 upper confidence limit.",
    },
    {
        "slide": 2,
        "question_en": "How is the nominal load of 0.98 calculated, and what does it imply?",
        "question_zh": "名义负荷0.98是怎样计算的，它意味着什么？",
        "answer_en": (
            "Urgent arrivals average 3.69 per day and routine arrivals average 2.19, giving "
            "5.88 arrivals against six appointment slots per day. Dividing 5.88 by six gives "
            "a nominal load of 0.98. This means the system has very little spare capacity, so "
            "prioritising one group mainly redistributes waiting time to another group."
        ),
        "answer_zh": (
            "紧急患者平均每天到达3.69人，常规患者平均每天到达2.19人，总计5.88人；"
            "每天只有6个预约时段。5.88除以6得到名义负荷0.98。这说明系统几乎没有闲置"
            "容量，因此优先照顾一类患者，主要会把等待时间转移给另一类患者。"
        ),
        "keywords": "0.98, arrival rates, six slots, spare capacity",
        "cue": "Number cue: (3.69 + 2.19) / 6 = 0.98.",
    },
    {
        "slide": 2,
        "question_en": "How did the literature shape the central trade-off?",
        "question_zh": "文献如何影响了核心权衡？",
        "answer_en": (
            "The literature suggested three relevant mechanisms: reserving capacity for "
            "higher-priority demand, changing priority as workload changes, and reusing "
            "released capacity through flexible rebooking. It also shows that these mechanisms "
            "have opportunity costs. That led me to treat the problem as a transparent trade-off "
            "between urgent access, routine access, and total system performance."
        ),
        "answer_zh": (
            "相关文献提出了三种有用机制：为高优先级需求预留容量、随着工作负荷改变优先级，"
            "以及通过灵活重排利用释放的容量。文献同时说明这些机制都有机会成本。因此，我把"
            "问题表述为紧急可及性、常规可及性和总体系统表现之间的透明权衡。"
        ),
        "keywords": "literature, capacity reservation, dynamic priority, rebooking",
        "cue": "Mechanism cue: protect, adapt, release and reuse.",
    },
    {
        "slide": 3,
        "question_en": "How do Q2, Q3 and Q4 compare?",
        "question_zh": "Q2、Q3和Q4的结果如何比较？",
        "answer_en": (
            "Q2 is the neutral benchmark, with urgent, routine and overall means all close "
            "to 4.148 days. Q3 strongly favours urgent access: 2.123 urgent, 9.053 routine and "
            "4.703 overall. Q4 is more balanced: 4.009 urgent, 4.574 routine and 4.219 overall. "
            "All three satisfy the overall five-day requirement, but they distribute waiting "
            "very differently."
        ),
        "answer_zh": (
            "Q2是中立基准，紧急、常规和总体平均等待都接近4.148天。Q3明显偏向紧急可及性："
            "紧急2.123天、常规9.053天、总体4.703天。Q4更加均衡：紧急4.009天、常规"
            "4.574天、总体4.219天。三者都满足总体5天要求，但等待时间的分配方式差异很大。"
        ),
        "keywords": "compare Q2 Q3 Q4, urgent, routine, overall",
        "cue": "Number cue: Q2 4.148; Q3 2.123/9.053; Q4 4.009/4.574.",
    },
    {
        "slide": 3,
        "question_en": "Why recommend Q3 despite its effect on routine patients?",
        "question_zh": "既然Q3会影响常规患者，为什么仍然推荐它？",
        "answer_en": (
            "The stated objective gives priority to urgent waiting while imposing an overall "
            "mean limit of five days. Under that objective, Q3 is the strongest feasible policy: "
            "urgent waiting is 2.123 days and the conservative overall upper confidence limit is "
            "4.893 days. I would not ignore the routine cost; I would add a routine-patient "
            "service guardrail before implementation."
        ),
        "answer_zh": (
            "题目设定的目标是在总体平均等待不超过5天的条件下优先降低紧急患者等待。按照这一"
            "目标，Q3是最强的可行策略：紧急等待为2.123天，总体均值的保守置信区间上限为"
            "4.893天。我不会忽略常规患者的代价；实际实施前应增加常规患者服务保障。"
        ),
        "keywords": "why Q3, objective, feasible, routine fairness",
        "cue": "Mechanism cue: best for the stated objective, conditional on a guardrail.",
    },
    {
        "slide": 4,
        "question_en": "Why use the same experiment design and stop at 50,000 arrivals?",
        "question_zh": "为什么三个策略使用相同实验设计，并在50,000次到达时停止？",
        "answer_en": (
            "Using the same warm-up, stopping rule, replication count and confidence method "
            "makes the policy comparison consistent. Stopping each replication at 50,000 "
            "arrivals provides a long post-warm-up observation period, which stabilises each "
            "replication mean while keeping computation manageable. Final precision is then "
            "judged across independent replications, not from run length alone."
        ),
        "answer_zh": (
            "三个策略采用相同的预热期、停止规则、重复次数和置信区间方法，可以保证比较口径"
            "一致。每次重复在50,000次到达时停止，能在预热后提供足够长的观测期，使单次重复"
            "的均值更稳定，同时保持计算量可控。最终精度仍通过独立重复之间的变异来判断。"
        ),
        "keywords": "same design, 50,000 arrivals, fair comparison, run length",
        "cue": "Number cue: one replication stops at 50,000 arrivals.",
    },
    {
        "slide": 4,
        "question_en": "Why use a warm-up of 1,000 served patients?",
        "question_zh": "为什么选择服务1,000名患者作为预热期？",
        "answer_en": (
            "The model starts from an empty artificial state, so early observations should "
            "not be included in steady-state estimates. A warm-up of 1,000 served patients "
            "was chosen after sensitivity checks at 500, 1,000 and 2,000. Their point estimates "
            "were close and their 95% confidence intervals overlapped, so the main conclusion "
            "was not sensitive to this reasonable choice."
        ),
        "answer_zh": (
            "模型从一个人为的空系统开始，因此早期观测不应进入稳态估计。选择服务1,000名"
            "患者作为预热期，是基于500、1,000和2,000三个设置的敏感性检查。它们的点估计"
            "相近，95%置信区间相互重叠，所以主要结论对这一合理选择并不敏感。"
        ),
        "keywords": "warm-up, empty start, 500 1000 2000, sensitivity",
        "cue": "Number cue: chosen warm-up = 1,000 served patients.",
    },
    {
        "slide": 4,
        "question_en": "Why use 250 independent replications and 95% confidence intervals?",
        "question_zh": "为什么使用250次独立重复和95%置信区间？",
        "answer_en": (
            "Independent replications provide an empirical sampling distribution for each "
            "reported mean. With 250 replications, every selected policy achieved the planned "
            "relative confidence-interval half-width below 5%. The 95% intervals show Monte "
            "Carlo uncertainty and also allow a conservative feasibility check by comparing "
            "the overall upper limit with five days."
        ),
        "answer_zh": (
            "独立重复可以形成每个报告均值的经验抽样分布。使用250次重复后，所有入选策略的"
            "相对置信区间半宽都低于预设的5%目标。95%置信区间展示了蒙特卡洛随机误差，也"
            "可以用总体区间上限与5天比较，进行较保守的可行性检查。"
        ),
        "keywords": "250 replications, 95% CI, precision, upper limit",
        "cue": "Number cue: 250 replications; relative half-width target below 5%.",
    },
    {
        "slide": 5,
        "question_en": "Why is FAS a useful benchmark?",
        "question_zh": "为什么最早可用时段策略是一个有用的基准？",
        "answer_en": (
            "First-available scheduling is simple, neutral between patient categories and easy "
            "to interpret. It assigns each patient the earliest open slot without explicit "
            "priority protection. That makes Q2 a clear baseline for measuring whether the "
            "extra mechanisms in Q3 or Q4 genuinely change urgent access and what cost they "
            "impose elsewhere."
        ),
        "answer_zh": (
            "最早可用时段策略简单、对两类患者保持中立，而且容易解释。它不设置明确的优先"
            "保护，而是把每位患者安排到最早的空闲时段。因此，Q2可以清楚衡量Q3或Q4新增"
            "机制是否真正改变了紧急可及性，以及这种改变在其他方面付出了什么代价。"
        ),
        "keywords": "FAS, benchmark, neutral, earliest slot",
        "cue": "Mechanism cue: no protection; assign the earliest open slot.",
    },
    {
        "slide": 5,
        "question_en": "How was the Q2 benchmark validated?",
        "question_zh": "Q2基准结果是怎样验证的？",
        "answer_en": (
            "The Q2 logic behaves as expected: with one merged first-available queue, urgent "
            "and routine means are almost identical. The overall mean is 4.148 days with a "
            "95% confidence interval from 3.978 to 4.318, and the relative half-width is about "
            "4.10%. The warm-up sensitivity intervals also overlap, supporting a stable and "
            "precise benchmark."
        ),
        "answer_zh": (
            "Q2的逻辑表现符合预期：在合并的最早可用规则下，紧急和常规患者的均值几乎相同。"
            "总体均值为4.148天，95%置信区间为3.978到4.318天，相对半宽约4.10%。不同"
            "预热期的置信区间也相互重叠，因此这一基准既稳定又具有足够精度。"
        ),
        "keywords": "validate Q2, equal class means, confidence interval, warm-up",
        "cue": "Number cue: 4.148 days; 95% CI [3.978, 4.318].",
    },
    {
        "slide": 6,
        "question_en": "How does Q3 work?",
        "question_zh": "Q3策略是怎样运行的？",
        "answer_en": (
            "Q3 observes the current urgent backlog, denoted by NU, and delays a selected "
            "routine appointment by round(alpha1 times NU plus alpha2 times NU squared). "
            "With alpha1 equal to zero and alpha2 equal to 0.35, protection grows nonlinearly "
            "as urgent congestion increases. This shifts near-term capacity from routine "
            "patients toward urgent patients."
        ),
        "answer_zh": (
            "Q3观察当前紧急积压量NU，并把选定的常规预约推迟round(alpha1乘以NU加alpha2"
            "乘以NU平方)个单位。当alpha1为0、alpha2为0.35时，保护强度会随着紧急拥堵"
            "以非线性方式增加，从而把近期容量从常规患者转向紧急患者。"
        ),
        "keywords": "how Q3 works, urgent backlog NU, nonlinear delay rule",
        "cue": "Mechanism cue: round(0.35 x NU squared).",
    },
    {
        "slide": 6,
        "question_en": "Why was alpha2 = 0.35 selected?",
        "question_zh": "为什么选择alpha2等于0.35？",
        "answer_en": (
            "Among the tested nonlinear protection settings, alpha2 equal to 0.35 gave the "
            "strongest urgent improvement while keeping the overall 95% confidence-interval "
            "upper limit below five days. It produced an urgent mean of 2.123 days and an "
            "overall upper limit of 4.893. Stronger protection pushed routine and overall "
            "performance beyond the feasible boundary."
        ),
        "answer_zh": (
            "在测试过的非线性保护设置中，alpha2等于0.35在总体95%置信区间上限仍低于5天"
            "的同时，带来了最强的紧急等待改善。其紧急均值为2.123天，总体区间上限为"
            "4.893天。更强的保护会使常规和总体表现越过可行边界。"
        ),
        "keywords": "why alpha2 0.35, parameter choice, feasibility boundary",
        "cue": "Number cue: alpha2 = 0.35; overall upper limit = 4.893.",
    },
    {
        "slide": 6,
        "question_en": "Why does routine waiting rise to 9.053 days?",
        "question_zh": "为什么常规患者等待会上升到9.053天？",
        "answer_en": (
            "Q3 does not add appointment slots. When the urgent backlog grows, it protects "
            "near-term capacity by moving selected routine appointments later. Because nominal "
            "load is already 0.98, there is very little spare capacity to absorb those delays. "
            "The 9.053-day routine mean is therefore the visible cost of reallocating scarce "
            "capacity toward urgent patients."
        ),
        "answer_zh": (
            "Q3并没有增加预约时段。当紧急积压上升时，它通过把选定的常规预约向后移动来保护"
            "近期容量。由于名义负荷已经是0.98，系统几乎没有多余容量来吸收这些延迟。因此，"
            "常规均值9.053天正是把稀缺容量转向紧急患者所产生的可见代价。"
        ),
        "keywords": "routine 9.053, fairness cost, reallocation, no new capacity",
        "cue": "Number cue: routine mean = 9.053 days.",
    },
    {
        "slide": 7,
        "question_en": "How does Q4 work, and what do eta, L and H mean?",
        "question_zh": "Q4怎样运行，eta、L和H分别表示什么？",
        "answer_en": (
            "Q4 protects near-term appointment capacity and then checks at the end of each day "
            "whether protected slots can be released. An eligible on-call routine patient with "
            "the longest wait is moved into a released slot. Eta equal to 0.2 controls the release "
            "threshold, while L equal to zero and H equal to three define the lower and upper "
            "boundaries of the protected horizon used in the policy."
        ),
        "answer_zh": (
            "Q4先保护近期预约容量，然后在每天结束时检查受保护时段是否可以释放。若有合适的"
            "空位，就把等待时间最长且符合条件的随叫随到常规患者提前安排。eta等于0.2控制"
            "释放阈值，L等于0和H等于3定义了该策略所使用保护区间的下界和上界。"
        ),
        "keywords": "how Q4 works, eta L H, hold release rebook",
        "cue": "Mechanism cue: hold, release, then rebook; eta 0.2, L 0, H 3.",
    },
    {
        "slide": 7,
        "question_en": "Why is Q4's urgent improvement small despite 97.911% utilisation?",
        "question_zh": "为什么Q4的利用率达到97.911%，紧急等待改善却很小？",
        "answer_en": (
            "High utilisation shows that released slots are not being wasted, and rebooked "
            "patients save an average of 0.515 days. However, utilisation is not the same as "
            "creating spare capacity. At a nominal load of 0.98, few protected slots remain "
            "truly empty, so Q4 can only reduce urgent waiting from 4.148 to 4.009 days."
        ),
        "answer_zh": (
            "高利用率说明释放的时段没有被大量浪费，被重新预约的患者平均节省0.515天。"
            "但是，利用率高并不等于创造了额外容量。在名义负荷0.98的情况下，真正保持空闲"
            "的受保护时段很少，因此Q4只能把紧急等待从4.148天下降到4.009天。"
        ),
        "keywords": "97.911% utilisation, small improvement, spare capacity, 0.515",
        "cue": "Number cue: 97.911% used, but urgent mean only improves to 4.009.",
    },
    {
        "slide": 8,
        "question_en": "Would Q3 be implemented unchanged?",
        "question_zh": "Q3可以不作修改就直接实施吗？",
        "answer_en": (
            "No. Q3 is the best answer to the stated optimisation objective, but a routine mean "
            "of 9.053 days is operationally and ethically important. Before implementation, I "
            "would add a routine-patient guardrail, such as a mean or percentile waiting-time "
            "limit, and then retune alpha2. I would also pilot the rule and monitor both patient "
            "groups rather than rely only on the overall mean."
        ),
        "answer_zh": (
            "不能。Q3是对题目优化目标的最佳回答，但常规患者平均等待9.053天在运营和伦理上"
            "都很重要。实施前，我会增加常规患者保障，例如平均等待或分位数等待上限，然后"
            "重新调整alpha2；还会先进行试点，同时监测两类患者，而不是只看总体均值。"
        ),
        "keywords": "implement Q3 unchanged, guardrail, retune, pilot",
        "cue": "Mechanism cue: add a routine service-level constraint before rollout.",
    },
    {
        "slide": 8,
        "question_en": "What are the main limitations and the next useful experiment?",
        "question_zh": "主要局限是什么，下一项有价值的实验是什么？",
        "answer_en": (
            "The model assumes stationary Poisson arrivals, deterministic appointment service, "
            "identical working days, and no cancellations, no-shows, overtime, preferences or "
            "clinical deterioration. The parameter search is also limited. The next useful "
            "experiment is a hybrid policy that combines Q3's workload signal with Q4's controlled "
            "release, while imposing an explicit routine-patient service guardrail."
        ),
        "answer_zh": (
            "模型假设到达过程是平稳泊松过程、预约服务时间确定、所有工作日相同，并且没有取消、"
            "爽约、加班、患者偏好或病情恶化；参数搜索范围也有限。下一项有价值的实验，是把"
            "Q3的工作负荷信号与Q4的受控释放机制结合成混合策略，同时加入明确的常规患者服务保障。"
        ),
        "keywords": "limitations, assumptions, next experiment, hybrid policy",
        "cue": "Mechanism cue: Q3 signal + Q4 release + routine guardrail.",
    },
]


def project_root() -> Path:
    configured = os.environ.get("MATH6185_PROJECT_ROOT")
    if configured:
        return Path(configured).resolve()
    candidate = Path(__file__).resolve().parents[2]
    if (candidate / "outputs").exists():
        return candidate
    return candidate.parents[1]


def output_path() -> Path:
    return (
        project_root()
        / "outputs"
        / "zw1f25_MATH6185_presentation"
        / "zw1f25_MATH6185_QA_Guide_PPT_Focused_Bilingual.docx"
    )


def validate_question_data() -> None:
    if len(QUESTIONS) != 18:
        raise ValueError(f"Expected 18 questions, found {len(QUESTIONS)}")
    expected_slides = [1, 1, 2, 2, 3, 3, 4, 4, 4, 5, 5, 6, 6, 6, 7, 7, 8, 8]
    actual_slides = [item["slide"] for item in QUESTIONS]
    if actual_slides != expected_slides:
        raise ValueError(f"Unexpected slide assignment: {actual_slides}")
    required = {
        "slide",
        "question_en",
        "question_zh",
        "answer_en",
        "answer_zh",
        "keywords",
        "cue",
    }
    for index, item in enumerate(QUESTIONS, start=1):
        if set(item) != required:
            raise ValueError(f"Q{index} fields do not match the schema")
        if any(not str(item[key]).strip() for key in required):
            raise ValueError(f"Q{index} contains an empty field")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def apply_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != 9360:
        raise ValueError(f"Table widths must total 9360 DXA: {widths}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, name="Calibri", size=11, bold=None, color=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    tokens = {
        "Heading 1": (16, "2E74B5", 18, 10),
        "Heading 2": (13, "2E74B5", 14, 7),
        "Heading 3": (12, "1F4D78", 10, 5),
    }
    for style_name, (size, color, before, after) in tokens.items():
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    custom_styles = {
        "Guide Title": (24, "17324D", True, 0, 8),
        "Chinese Question": (11, "374151", True, 0, 5),
        "Answer Chinese": (10, "374151", False, 0, 7),
        "Listening Cue": (9, "1F6F8B", False, 0, 3),
        "Memory Cue": (9, "7A5A00", True, 0, 10),
        "Guide Note": (10, "687783", False, 0, 10),
    }
    for name, (size, color, bold, before, after) in custom_styles.items():
        if name not in styles:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = styles[name]
        style.base_style = normal
        style.font.name = "Calibri"
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_end])


def add_labelled_paragraph(doc: Document, label: str, text: str, style=None) -> None:
    paragraph = doc.add_paragraph(style=style)
    label_run = paragraph.add_run(label)
    label_run.bold = True
    label_run.font.color.rgb = RGBColor.from_string("17324D")
    text_run = paragraph.add_run(text)
    text_run._element.get_or_add_rPr().get_or_add_rFonts().set(
        qn("w:eastAsia"), "Microsoft YaHei"
    )


def build_document(target: Path) -> None:
    validate_question_data()
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header.paragraphs[0]
    header.text = "MATH6185 | PPT-Focused Q&A"
    set_run_font(header.runs[0], size=9, color="687783")
    footer = section.footer.paragraphs[0]
    add_page_field(footer)

    title = doc.add_paragraph(style="Guide Title")
    title.add_run("MATH6185 PPT-Focused\nBilingual Q&A Guide")
    subtitle = doc.add_paragraph(style="Guide Note")
    subtitle.add_run(
        "Slides 1-8 only | English answers with Chinese study support | Student ID: zw1f25"
    )
    note = doc.add_paragraph(style="Guide Note")
    note.add_run(
        "Use: listen for one keyword, locate the question number, answer directly, "
        "then give the number or mechanism cue. Chinese text is for preparation, "
        "not for reading aloud in the English presentation."
    )

    doc.add_paragraph("Quick location", style="Heading 1")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headings = ("Slide", "Listen for", "Questions")
    for cell, text in zip(table.rows[0].cells, headings):
        cell.text = text
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=9, bold=True, color="17324D")
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "E8EEF5")
        cell._tc.get_or_add_tcPr().append(shading)
    set_repeat_table_header(table.rows[0])

    quick_rows = [
        ("Slide 1", "decision, main result, why Q3", "Q1-Q2"),
        ("Slide 2", "0.98 load, literature, trade-off", "Q3-Q4"),
        ("Slide 3", "compare policies, fairness", "Q5-Q6"),
        ("Slide 4", "50,000, warm-up, 250, 95% CI", "Q7-Q9"),
        ("Slide 5", "FAS, benchmark, validation", "Q10-Q11"),
        ("Slide 6", "Q3 rule, alpha2, routine 9.053", "Q12-Q14"),
        ("Slide 7", "Q4, eta L H, utilisation", "Q15-Q16"),
        ("Slide 8", "implementation, limits, next step", "Q17-Q18"),
    ]
    for slide, listen, questions in quick_rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, (slide, listen, questions)):
            cell.text = text
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_run_font(run, size=9)
        cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    apply_table_geometry(table, [1200, 6040, 2120])

    doc.add_page_break()
    by_slide = {
        1: "Decision and result",
        2: "Capacity pressure and literature framing",
        3: "Cross-policy comparison",
        4: "Experimental design",
        5: "Q2 benchmark",
        6: "Q3 workload-responsive policy",
        7: "Q4 OCR policy",
        8: "Recommendation, limitations and next work",
    }
    for slide_number, section_title in by_slide.items():
        doc.add_paragraph(f"Slide {slide_number} - {section_title}", style="Heading 1")
        for index, item in enumerate(QUESTIONS, start=1):
            if item["slide"] != slide_number:
                continue
            heading = doc.add_paragraph(
                f"Q{index}. {item['question_en']}", style="Heading 2"
            )
            heading.paragraph_format.keep_with_next = True
            chinese_question = doc.add_paragraph(
                f"问题中文：{item['question_zh']}", style="Chinese Question"
            )
            chinese_question.paragraph_format.keep_with_next = True
            add_labelled_paragraph(
                doc, "Answer in English: ", item["answer_en"]
            )
            add_labelled_paragraph(
                doc,
                "中文对照：",
                item["answer_zh"],
                style="Answer Chinese",
            )
            doc.add_paragraph(
                f"Listening keywords / 听题关键词：{item['keywords']}",
                style="Listening Cue",
            )
            doc.add_paragraph(
                f"Memory cue / 记忆提示：{item['cue']}",
                style="Memory Cue",
            )

    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)


def main() -> None:
    target = output_path()
    build_document(target)
    print(target)


if __name__ == "__main__":
    main()
