from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
OUTPUT = ROOT / "outputs" / "zw1f25_MATH6185_presentation"
CONTENT_PATH = HERE / "presentation_content.json"
SCRIPT_PATH = OUTPUT / "zw1f25_MATH6185_Speaker_Script.docx"
QA_PATH = OUTPUT / "zw1f25_MATH6185_QA_Guide.docx"

NAVY = "17324D"
TEAL = "007F82"
PALE_TEAL = "DCEEEF"
INK = "18222C"
GREY = "687783"
LIGHT_GREY = "EEF2F4"
AMBER = "D68B00"
PALE_AMBER = "FFF4D6"
LINE = "CDD7DC"


def set_run_font(run, name: str, size: float, color: str = INK, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_border(paragraph, side: str, color: str, size: int = 10) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    edge = p_bdr.find(qn(f"w:{side}"))
    if edge is None:
        edge = OxmlElement(f"w:{side}")
        p_bdr.append(edge)
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), str(size))
    edge.set(qn("w:space"), "6")
    edge.set(qn("w:color"), color)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_end])
    set_run_font(run, "Aptos", 9, GREY)


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(NAVY)
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)

    h1 = styles["Heading 1"]
    h1.font.name = "Aptos Display"
    h1.font.size = Pt(18)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor.from_string(NAVY)
    h1._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    h1._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    h1.paragraph_format.keep_with_next = True

    h2 = styles["Heading 2"]
    h2.font.name = "Aptos"
    h2.font.size = Pt(13)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor.from_string(TEAL)
    h2._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    h2._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(4)
    h2.paragraph_format.keep_with_next = True

    if "Muted note" not in [style.name for style in styles]:
        muted = styles.add_style("Muted note", WD_STYLE_TYPE.PARAGRAPH)
    else:
        muted = styles["Muted note"]
    muted.base_style = normal
    muted.font.name = "Aptos"
    muted.font.size = Pt(10)
    muted.font.italic = True
    muted.font.color.rgb = RGBColor.from_string(GREY)
    muted.paragraph_format.left_indent = Cm(0.5)
    muted.paragraph_format.space_before = Pt(3)
    muted.paragraph_format.space_after = Pt(5)

    if "Evidence" not in [style.name for style in styles]:
        evidence = styles.add_style("Evidence", WD_STYLE_TYPE.PARAGRAPH)
    else:
        evidence = styles["Evidence"]
    evidence.base_style = normal
    evidence.font.name = "Aptos"
    evidence.font.size = Pt(9.5)
    evidence.font.italic = True
    evidence.font.color.rgb = RGBColor.from_string(GREY)
    evidence.paragraph_format.space_before = Pt(3)
    evidence.paragraph_format.space_after = Pt(12)


def configure_document(document: Document, title: str) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.9)

    configure_styles(document)
    props = document.core_properties
    props.title = title
    props.author = "zw1f25"
    props.subject = "MATH6185 Case Study 1 oral presentation preparation"
    props.keywords = "MATH6185, simulation, oral presentation, zw1f25"

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header_paragraph.add_run("MATH6185 | CASE STUDY 1 | zw1f25")
    set_run_font(header_run, "Aptos", 9, GREY, True)
    set_paragraph_border(header_paragraph, "bottom", LINE, 6)

    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Cm(17))
    table.autofit = False
    table.columns[0].width = Cm(13)
    table.columns[1].width = Cm(4)
    left_cell = table.cell(0, 0)
    right_cell = table.cell(0, 1)
    left_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    right_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    left_p = left_cell.paragraphs[0]
    right_p = right_cell.paragraphs[0]
    left_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    right_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    left_run = left_p.add_run(title)
    set_run_font(left_run, "Aptos", 9, GREY)
    add_page_number(right_p)


def add_cover_title(document: Document, title: str, subtitle: str) -> None:
    paragraph = document.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(18)
    paragraph.add_run(title)
    set_paragraph_border(paragraph, "bottom", TEAL, 14)
    sub = document.add_paragraph()
    sub.paragraph_format.space_after = Pt(18)
    run = sub.add_run(subtitle)
    set_run_font(run, "Aptos", 13, GREY)


def add_key_message(document: Document, heading: str, text: str, risk: bool = False) -> None:
    table = document.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Cm(17)
    cell = table.cell(0, 0)
    set_cell_shading(cell, PALE_AMBER if risk else PALE_TEAL)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    heading_run = paragraph.add_run(f"{heading}\n")
    set_run_font(heading_run, "Aptos", 12, AMBER if risk else TEAL, True)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, "Aptos", 11, NAVY, True)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_script_slide(document: Document, slide: dict) -> None:
    heading = document.add_paragraph(
        f"Slide {slide['number']} - {slide['title']}", style="Heading 1"
    )
    heading.paragraph_format.page_break_before = True
    heading.paragraph_format.keep_with_next = True

    purpose = document.add_paragraph()
    purpose.paragraph_format.keep_with_next = True
    purpose_run = purpose.add_run(f"Purpose: {slide['purpose']}")
    set_run_font(purpose_run, "Aptos", 10.5, TEAL, True)

    section_heading = document.add_paragraph("Suggested wording", style="Heading 2")
    section_heading.paragraph_format.keep_with_next = True
    notes = document.add_paragraph(slide["speakerNotes"])
    notes.paragraph_format.keep_together = True

    if slide["shortVersionOmissions"]:
        document.add_paragraph("Shorter version", style="Heading 2")
        for omission in slide["shortVersionOmissions"]:
            paragraph = document.add_paragraph(style="Muted note")
            paragraph.add_run(f"[Optional if time is short] {omission}")

    sources = document.add_paragraph(style="Evidence")
    sources.add_run("Evidence: " + " ".join(slide["sources"]))


def build_speaker_script(content: dict, output_path: Path) -> None:
    document = Document()
    configure_document(document, "MATH6185 Speaker Script")
    add_cover_title(
        document,
        "MATH6185 Oral Presentation - Speaker Script",
        "Student ID: zw1f25 | Result-first 8-12 minute version",
    )
    add_key_message(
        document,
        "Core message",
        "Recommend Q3: urgent mean waiting falls from 4.148 to 2.123 days, while the overall upper 95% confidence limit remains below five days.",
    )
    add_key_message(
        document,
        "Critical qualification",
        "Routine mean waiting rises to 9.053 days. Present the recommendation as conditional on a routine service-level guardrail.",
        risk=True,
    )

    document.add_paragraph("Delivery guide", style="Heading 1")
    delivery_points = [
        ("Structure", "Slides 1-8 are the spoken core; slides 9-14 are backup for questions."),
        ("Timing", "Aim for roughly 60-75 seconds on slides 1-4 and 45-60 seconds on slides 5-8."),
        ("Short version", "If time is limited, omit the grey optional lines rather than rushing the final recommendation."),
        ("Demonstration", "Do not run AnyLogic live. Use the model-structure backup slide if implementation details are requested."),
        ("Questions", "Answer directly, give one result or mechanism, then state the limitation if it matters."),
    ]
    for label, text in delivery_points:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(5)
        label_run = paragraph.add_run(f"{label}: ")
        set_run_font(label_run, "Aptos", 11, NAVY, True)
        text_run = paragraph.add_run(text)
        set_run_font(text_run, "Aptos", 11, INK)

    for slide in [item for item in content["slides"] if item["kind"] == "core"]:
        add_script_slide(document, slide)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def add_qa_entry(document: Document, entry: dict, number: int) -> None:
    question = document.add_paragraph(f"Q{number}. {entry['question']}", style="Heading 1")
    question.paragraph_format.keep_with_next = True
    if number == 21:
        question.paragraph_format.page_break_before = True

    answer = document.add_paragraph()
    answer.paragraph_format.keep_with_next = True
    label = answer.add_run("Answer in English: ")
    set_run_font(label, "Aptos", 11, TEAL, True)
    body = answer.add_run(entry["answerEnglish"])
    set_run_font(body, "Aptos", 11, INK)

    chinese = document.add_paragraph()
    chinese.paragraph_format.left_indent = Cm(0.35)
    chinese.paragraph_format.right_indent = Cm(0.35)
    chinese.paragraph_format.space_before = Pt(4)
    chinese.paragraph_format.space_after = Pt(5)
    chinese.paragraph_format.keep_with_next = True
    set_paragraph_shading(chinese, PALE_TEAL)
    chinese_label = chinese.add_run("中文理解：")
    set_run_font(chinese_label, "Microsoft YaHei", 10.5, NAVY, True)
    chinese_body = chinese.add_run(entry["explanationChinese"])
    set_run_font(chinese_body, "Microsoft YaHei", 10.5, INK)

    evidence = document.add_paragraph(style="Evidence")
    evidence.paragraph_format.keep_together = True
    evidence.add_run("Evidence: " + entry["evidence"])


def build_qa_guide(content: dict, output_path: Path) -> None:
    document = Document()
    configure_document(document, "MATH6185 Q&A Guide")
    add_cover_title(
        document,
        "MATH6185 Oral Presentation - Q&A Guide",
        f"Student ID: zw1f25 | {len(content['qa'])} likely questions with English answers and Chinese explanations",
    )
    add_key_message(
        document,
        "Answer pattern",
        "Lead with the conclusion, support it with one number or model mechanism, then state the relevant limitation. Do not overclaim global optimality or direct AnyLogic validation.",
    )
    document.add_paragraph("High-priority facts to memorise", style="Heading 1")
    facts = [
        ("Q2 FAS", "Overall mean 4.148 days."),
        ("Q3", "Urgent 2.123, routine 9.053, overall 4.703 days; urgent reduction 48.8%."),
        ("Q4", "Urgent 4.009, routine 4.574, overall 4.219 days; utilisation 97.911%."),
        ("Experiment", "1,000 served warm-up, 50,000-arrival stop, 250 independent replications, 95% confidence intervals."),
        ("Recommendation", "Q3 for the stated objective, with a routine service-level guardrail."),
    ]
    for label, fact in facts:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.2)
        label_run = paragraph.add_run(f"{label}: ")
        set_run_font(label_run, "Aptos", 11, TEAL, True)
        run = paragraph.add_run(fact)
        set_run_font(run, "Aptos", 11, INK)

    for number, entry in enumerate(content["qa"], start=1):
        add_qa_entry(document, entry, number)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> None:
    content = json.loads(CONTENT_PATH.read_text(encoding="utf-8"))
    build_speaker_script(content, SCRIPT_PATH)
    build_qa_guide(content, QA_PATH)
    print(f"HANDOUT BUILD PASS: {SCRIPT_PATH}")
    print(f"HANDOUT BUILD PASS: {QA_PATH}")


if __name__ == "__main__":
    main()
