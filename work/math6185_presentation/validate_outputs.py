from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader


HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
ORIGINAL_ROOT = ROOT.parent.parent if ROOT.parent.name == ".worktrees" else ROOT
OUT = ROOT / "outputs" / "zw1f25_MATH6185_presentation"
PPTX = OUT / "zw1f25_MATH6185_Presentation.pptx"
PDF = OUT / "zw1f25_MATH6185_Presentation.pdf"
SCRIPT = OUT / "zw1f25_MATH6185_Speaker_Script.docx"
QA = OUT / "zw1f25_MATH6185_QA_Guide.docx"
DATA = json.loads((HERE / "source_data.json").read_text(encoding="utf-8"))


def document_text(document: Document) -> str:
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            paragraphs.extend(cell.text for cell in row.cells)
    return "\n".join(paragraphs)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


for path in (PPTX, PDF, SCRIPT, QA):
    assert path.exists() and path.stat().st_size > 20_000, path

with ZipFile(PPTX) as archive:
    names = archive.namelist()
    slide_files = [
        name
        for name in names
        if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
    ]
    notes_files = [
        name
        for name in names
        if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", name)
    ]
    assert len(slide_files) == 14, len(slide_files)
    assert len(notes_files) == 14, len(notes_files)
    evidence_xml = [
        name
        for name in names
        if name.endswith(".xml")
        and (
            name.startswith("ppt/slides/")
            or name.startswith("ppt/charts/")
            or name.startswith("ppt/notesSlides/")
        )
    ]
    ppt_text = " ".join(
        archive.read(name).decode("utf-8", errors="ignore") for name in evidence_xml
    )

pdf = PdfReader(str(PDF))
assert len(pdf.pages) == 14, len(pdf.pages)
pdf_text = "\n".join(page.extract_text() or "" for page in pdf.pages)

script_doc = Document(SCRIPT)
qa_doc = Document(QA)
script_text = document_text(script_doc)
qa_text = document_text(qa_doc)

for section in (script_doc.sections[0], qa_doc.sections[0]):
    assert abs(section.page_width.cm - 21.0) < 0.05
    assert abs(section.page_height.cm - 29.7) < 0.05
    for margin in (
        section.top_margin,
        section.bottom_margin,
        section.left_margin,
        section.right_margin,
    ):
        assert abs(margin.cm - 2.0) < 0.05

required_tokens = (
    "4.148",
    "2.123",
    "9.053",
    "4.703",
    "4.009",
    "4.574",
    "4.219",
    "97.911",
    "0.515",
)
for token in required_tokens:
    assert token in ppt_text, f"{token} missing from PPTX"
    assert token in script_text or token in qa_text, f"{token} missing from handouts"

for title in (
    "Reducing urgent waits",
    "Recommend Q3",
    "Full 95% confidence intervals",
):
    assert title in pdf_text, f"{title} missing from PDF"

assert "250 independent replications" in script_text
assert "50,000 arrivals" in script_text
assert "中文理解" in qa_text
assert "AI assistance" in qa_text
assert len(re.findall(r"(?m)^Q\d+\.", qa_text)) == 22
assert len(re.findall(r"(?m)^Slide \d+ -", script_text)) == 8

for forbidden in ("{{", "}}", "TODO", "PLACEHOLDER"):
    assert forbidden not in ppt_text
    assert forbidden not in script_text
    assert forbidden not in qa_text

assert DATA["q2"]["overallMean"] == 4.147829551134296
assert DATA["q3"]["alpha2"] == 0.35
assert DATA["q3"]["overallCI"][1] < DATA["system"]["overallMeanConstraintDays"]
assert DATA["q4"]["eta"] == 0.2
assert DATA["q4"]["H"] == 3

submission = ORIGINAL_ROOT / "outputs" / "zw1f25_MATH6185_submission"
expected_hashes = {
    "zw1f25_MATH6185_Report.pdf": "7DD574347034EF6B055C1AD60908C7D1BA828D6898A7E53991A1B8B0BAEA942C",
    "zw1f25_MATH6185_Q2_FAS.alp": "AA8F471BECA3C1AAF987561FBA1EA259CCBFC3484BB08E733C126995270E2C76",
    "zw1f25_MATH6185_Q3_WorkloadPolicy_alpha035.alp": "21AB826240270C02A8B703F55B92ECA402798F5E535FAA73FC9C6EA90E086BFE",
    "zw1f25_MATH6185_Q4_OCR_eta020_L0_H3.alp": "78C83C924107C12B98E9C85834138FFA4E63AEBF0E79D26F1D8EA31C8FDB830F",
}
for filename, expected in expected_hashes.items():
    source = submission / filename
    assert source.exists(), source
    assert sha256(source) == expected, filename

print("FINAL OUTPUT VALIDATION PASS")
